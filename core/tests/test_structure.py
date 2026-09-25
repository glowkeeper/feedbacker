"""Inspection reports structure only: never text, metadata values, or names."""

from __future__ import annotations

import json

import pytest
from helpers import PACK, make_zip

from feedbacker_core import cli
from feedbacker_core.structure import InspectionError, inspect_path, name_shape

SEEDED = json.loads((PACK / "seeded-identifiers.json").read_text())


def all_identifiers():
    for ids in SEEDED.values():
        yield ids["metadata_author"]
        for key in ("names", "student_ids", "emails", "urls", "organisations"):
            yield from ids[key]


def assert_no_content(lines):
    joined = "\n".join(lines)
    for value in all_identifiers():
        assert value not in joined, value
    for word in ("Plant Swap", "MoSCoW", "WebSockets", "Good choice of framework", "Grade: 62"):
        assert word not in joined


def test_name_shape_hides_letters_and_digits():
    assert name_shape("Pike_Jordan_100200302_final Report.DOCX") == (
        "Aaaa_Aaaaaa_999999999_aaaaa Aaaaaa.docx"
    )


def test_inspect_docx_and_pdf_show_no_content():
    for f in ("sub-a.docx", "sub-c.docx", "sub-b.pdf", "sub-d.pdf"):
        lines = inspect_path(PACK / "submissions" / f)
        assert lines[0].startswith("type:")
        assert_no_content(lines)


def test_inspect_marked_view_reveals_its_structure_only():
    lines = inspect_path(PACK / "marked-view-replica.pdf")
    assert "pages: 7" in lines[0]
    page = {ln.split(":")[0]: ln for ln in lines[2:]}
    assert all("largest-image=89%" in page[p] for p in ("p2", "p3", "p4"))
    assert "comment-headings=3" in page["p5"]
    assert "band-labels=24" in page["p6"]
    assert_no_content(lines)


def test_inspect_zip_shows_name_shapes_only(tmp_path):
    z = make_zip(
        tmp_path / "o.zip",
        {
            "Pike_Jordan_100200302_report.pdf": b"x",
            "Quill_Avery_100200301_report.docx": b"x",
        },
    )
    lines = inspect_path(z)
    assert "files: 2" in lines[0]
    assert any("Aaaa_Aaaaaa_999999999_aaaaaa.pdf" in ln for ln in lines)
    assert_no_content(lines)
    assert not any("100200302" in ln for ln in lines)


def test_cli_inspect(capsys):
    assert cli.main(["inspect", str(PACK / "submissions" / "sub-b.pdf")]) == 0
    out = capsys.readouterr().out
    assert "type: pdf" in out
    assert_no_content(out.splitlines())


# --- Review fixes: non-ASCII names, style names, damaged files ----------------


def test_name_shape_masks_any_script():
    assert name_shape("García_Élodie_Ōtsuka_12345.docx") == "Aaaaaa_Aaaaaa_Aaaaaa_99999.docx"
    assert name_shape("Иван Петров.pdf") == "Aaaa Aaaaaa.pdf"
    assert name_shape("名前 レポート.docx") == "aa aaaa.docx"
    assert name_shape("report.Pérez") == "aaaaaa.Aaaaa"  # a non-ASCII "extension" is masked too


def test_docx_style_names_are_bucketed(tmp_path):
    from docx import Document
    from docx.enum.style import WD_STYLE_TYPE

    doc = Document()
    doc.styles.add_style("Quill Avery Notes", WD_STYLE_TYPE.PARAGRAPH)
    doc.add_paragraph("x", style="Quill Avery Notes")
    doc.add_heading("y", level=1)
    path = tmp_path / "s.docx"
    doc.save(path)
    joined = "\n".join(inspect_path(path))
    assert "Quill" not in joined
    assert "paragraph style kinds: {'other': 1, 'heading': 1}" in joined


@pytest.mark.parametrize("name", ["bad.pdf", "bad.docx", "bad.zip"])
def test_damaged_files_fail_cleanly(tmp_path, name, capsys):
    path = tmp_path / name
    path.write_bytes(b"Quill Avery secret content")
    with pytest.raises(InspectionError) as err:
        inspect_path(path)
    assert "Quill" not in str(err.value) and "damaged" in str(err.value)
    assert cli.main(["inspect", str(path)]) == 1
    assert "could not inspect" in capsys.readouterr().err


def test_unsupported_type(tmp_path):
    path = tmp_path / "x.odt"
    path.write_bytes(b"")
    with pytest.raises(InspectionError, match="unsupported file type"):
        inspect_path(path)

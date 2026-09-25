"""Generate the synthetic fixture pack.

Every name, identifier, organisation, and piece of work here is fictional.
Emails and URLs use reserved example domains (RFC 2606).

Usage (from ``core/``):

    uv run --group fixtures python scripts/make_fixtures.py
"""

from __future__ import annotations

import csv
import hashlib
import io
import json
import zipfile
from datetime import UTC, datetime
from pathlib import Path

from docx import Document
from PIL import Image, ImageDraw
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.utils import ImageReader
from reportlab.pdfgen import canvas
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

from feedbacker_core.models import (
    Actor,
    ActorKind,
    Annotation,
    BandCount,
    Criterion,
    ImportRoute,
    JudgementEntry,
    Level,
    ModerationContext,
    ModerationRecord,
    ModeratorJudgement,
    OriginalAssessment,
    OriginalCriterionMark,
    Provenance,
    ReviewMode,
    Rubric,
    SourceFormat,
    SourceKind,
    Submission,
    SubmissionVerdict,
    Transformation,
    Verdict,
)

PACK = Path(__file__).resolve().parents[2] / "fixtures" / "synthetic" / "pack-01"
FIXED = datetime(2026, 1, 15, 9, 0, tzinfo=UTC)
MODERATOR = Actor(kind=ActorKind.MODERATOR, label="moderator")
MARKER = Actor(kind=ActorKind.ORIGINAL_MARKER, label="original marker")


def at(minutes: int) -> datetime:
    return FIXED.replace(minute=minutes)


# --- Rubric -----------------------------------------------------------------

# Points-based levels, in the style of Turnitin/Canvas rubrics: many levels,
# each with one points value and a label carrying its classification.
POINTS = [85, 75, 68, 62, 55, 48, 42, 35, 20]
BANDS = [(70, "1ST", 4), (60, "2:1", 3), (50, "2:2", 2), (40, "3RD", 1), (0, "FAIL", 0)]
# One deliberately inconsistent label, as found in real rubrics: 68 falls in
# the 2:1 band but is labelled 2:2. Importers must keep labels exactly as written.
MISLABELLED = {("testing", 68): "2:2 (68)"}


def band(points: int) -> tuple[str, int]:
    return next((name, idx) for floor, name, idx in BANDS if points >= floor)


CRITERIA = {
    "design": (
        "Requirements and design",
        "Identifies requirements and justifies an appropriate design.",
        [
            "Requirements missing or unrelated to the brief.",
            "Some requirements listed; design largely undocumented.",
            "Requirements listed; design described with limited justification.",
            "Clear requirements; design justified with appropriate diagrams.",
            "Thorough, prioritised requirements; design critically justified against alternatives.",
        ],
    ),
    "implementation": (
        "Implementation",
        "Builds a working application that meets the requirements.",
        [
            "Application does not run or bears little relation to the design.",
            "Partially working; significant requirements unmet.",
            "Working core features; code quality inconsistent.",
            "Most requirements met; well-structured, readable code.",
            "All requirements met; robust, well-structured code with thoughtful extras.",
        ],
    ),
    "testing": (
        "Testing and evaluation",
        "Tests the application and evaluates it against the requirements.",
        [
            "No meaningful testing.",
            "Ad hoc manual testing only.",
            "Planned testing of main features; limited evaluation.",
            "Systematic testing, including automated tests; evaluation against requirements.",
            "Comprehensive testing strategy; critical evaluation with evidence.",
        ],
    ),
    "reflection": (
        "Reflection and professional practice",
        "Reflects on the process and relates it to professional practice.",
        [
            "No reflection.",
            "Descriptive account of what was done.",
            "Some reflection on challenges, with limited insight.",
            "Insightful reflection linked to professional practice.",
            "Critical, well-evidenced reflection with clear actions for future practice.",
        ],
    ),
}


def build_rubric() -> Rubric:
    criteria = []
    for cid, (title, desc, descriptors) in CRITERIA.items():
        levels = []
        for p in POINTS:
            name, idx = band(p)
            label = MISLABELLED.get((cid, p), f"{name} ({p})")
            levels.append(Level(id=f"p{p}", label=label, descriptor=descriptors[idx], points=p))
        criteria.append(
            Criterion(
                id=cid, title=title, description=desc, weight=25, max_points=100, levels=levels
            )
        )
    return Rubric(
        id="cmp-5001-cw1",
        version="1.0",
        title="CMP5001 Web Application Development: Coursework 1 (synthetic)",
        criteria=criteria,
        provenance=Provenance(
            source="synthetic fixture generator",
            transformation=Transformation.IMPORTED,
            actor=MODERATOR,
            timestamp=FIXED,
        ),
    )


# --- Submissions --------------------------------------------------------------
# Each report seeds fictional direct and indirect identifiers for #16.

REPORTS = [
    {
        "id": "sub-a",
        "pseudonym": "[STUDENT_A]",
        "format": SourceFormat.DOCX,
        "author": "Avery Quill",
        "identifiers": {
            "names": ["Avery Quill"],
            "student_ids": ["S0000101"],
            "emails": ["avery.quill@example.com"],
            "urls": ["https://example.org/aquill/portfolio"],
            "organisations": ["Northwind Widgets Ltd"],
        },
        "title": "Plant Swap: a community plant-sharing web app",
        "sections": [
            ("Cover", "Student: Avery Quill (S0000101). Contact: avery.quill@example.com."),
            (
                "Requirements and design",
                "I gathered requirements by interviewing four members of a local gardening group "
                "and prioritised them using MoSCoW. The must-haves were listing a plant, browsing "
                "listings by postcode area, and messaging another member. I compared a server-rendered "
                "design with a single-page application and chose server rendering because most users "
                "browse on older phones. The entity-relationship diagram shows Member, Listing, and "
                "Message, and I explain why messages are not deleted when a listing is removed.",
            ),
            (
                "Implementation",
                "The application uses Flask with SQLite. All must-have and two should-have requirements "
                "are implemented, including image upload with server-side resizing. Routes are grouped "
                "into blueprints and database access is isolated in a repository module. Input is "
                "validated on the server and CSRF protection is enabled.",
            ),
            (
                "Testing and evaluation",
                "I wrote 42 pytest tests covering models, routes, and validation, reaching 87% line "
                "coverage, and ran a usability session with three group members. The evaluation table "
                "maps each requirement to its tests and to usability findings, and notes that search "
                "by postcode area was slower than expected with 5,000 synthetic listings.",
            ),
            (
                "Reflection",
                "My part-time role at Northwind Widgets Ltd taught me to write tickets before coding, "
                "which I applied here. I underestimated image handling and would spike risky features "
                "earlier next time. My code is at https://example.org/aquill/portfolio.",
            ),
        ],
    },
    {
        "id": "sub-b",
        "pseudonym": "[STUDENT_B]",
        "format": SourceFormat.PDF,
        "author": "Jordan Pike",
        "identifiers": {
            "names": ["Jordan Pike"],
            "student_ids": ["S0000102"],
            "emails": ["j.pike@example.com"],
            "urls": [],
            "organisations": [],
        },
        "title": "Study Buddy: a revision timetable planner",
        "sections": [
            ("Cover", "Jordan Pike, student number S0000102, j.pike@example.com."),
            (
                "Requirements and design",
                "The app should let students add modules and exam dates and get a timetable. "
                "I used a single-page application because it is modern.",
            ),
            (
                "Implementation",
                "I built the front end in React. Adding modules and exam dates works. The timetable "
                "generator sometimes schedules sessions after the exam date and I did not have time to "
                "fix this. Data is stored in the browser.",
            ),
            (
                "Testing and evaluation",
                "I tested the app by clicking through it and it mostly worked.",
            ),
            ("Reflection", "I learned React. Next time I would start earlier."),
        ],
    },
    {
        "id": "sub-c",
        "pseudonym": "[STUDENT_C]",
        "format": SourceFormat.DOCX,
        "author": "Riley Marsh",
        "identifiers": {
            "names": ["Riley Marsh", "Sam"],
            "student_ids": ["S0000103"],
            "emails": ["riley.marsh@example.com"],
            "urls": ["https://example.net/rmarsh"],
            "organisations": [],
        },
        "title": "Bike Fix Log: tracking repairs for a cycling club",
        "sections": [
            ("Cover", "Riley Marsh | S0000103 | riley.marsh@example.com"),
            (
                "Requirements and design",
                "Requirements came from the club's repair volunteer, Sam. I documented user stories "
                "and a use-case diagram. The design uses a REST API with a small JavaScript front end; "
                "the justification is brief.",
            ),
            (
                "Implementation",
                "Core features work: logging a repair, listing repairs per bike, and a parts total. "
                "Some route handlers mix database queries with formatting and there is duplicated "
                "validation code. Authentication is not implemented.",
            ),
            (
                "Testing and evaluation",
                "I wrote a test plan with twelve manual test cases and eight unit tests for the parts "
                "calculation. Two test cases failed and are listed as known issues. The evaluation "
                "briefly compares the result with the user stories.",
            ),
            (
                "Reflection",
                "Working with a real user was useful, although I found it hard to say no to extra "
                "requests. My screenshots on https://example.net/rmarsh show the progress.",
            ),
        ],
    },
    {
        "id": "sub-d",
        "pseudonym": "[STUDENT_D]",
        "format": SourceFormat.PDF,
        "author": "Casey Rowan",
        "identifiers": {
            "names": ["Casey Rowan"],
            "student_ids": ["S0000104"],
            "emails": ["casey.rowan@example.com"],
            "urls": [],
            "organisations": ["Fabrikam Games"],
        },
        "title": "Quiz Night: a real-time pub quiz app",
        "sections": [
            ("Cover", "Name: Casey Rowan. ID: S0000104. Email: casey.rowan@example.com."),
            (
                "Requirements and design",
                "I analysed three existing quiz apps and derived requirements, prioritised by risk. "
                "The design uses WebSockets for real-time scoring; I justify this against polling with "
                "a latency estimate and include sequence diagrams for joining a game and answering.",
            ),
            (
                "Implementation",
                "Implemented in Node.js with Socket.IO and PostgreSQL. All requirements are met, plus "
                "reconnection handling for dropped players. The code is modular with a clear "
                "separation between game logic and transport.",
            ),
            (
                "Testing and evaluation",
                "Unit tests cover the game state machine (31 tests) and an automated load test ran "
                "200 simulated players. I evaluate latency against my design estimate and discuss "
                "where it fell short.",
            ),
            (
                "Reflection",
                "My summer placement at Fabrikam Games shaped how I planned sprints. I reflect on a "
                "poor early decision to store state in memory and how refactoring it cost two weeks, "
                "and I set out what I will do differently.",
            ),
        ],
    },
]

# Original marker's points per criterion, summary comment, and inline comments,
# as they would appear in a Turnitin "current view". sub-b is marked generously
# on purpose so the moderation has something to find.
ORIGINAL = {
    "sub-a": (
        {"design": 85, "implementation": 85, "testing": 75, "reflection": 68},
        "Excellent, well-evidenced work throughout.",
        [("Strong use of MoSCoW prioritisation.", "prioritised them using MoSCoW", 1)],
    ),
    "sub-b": (
        {"design": 68, "implementation": 68, "testing": 55, "reflection": 55},
        "Good app, nice use of React.",
        [
            (
                "Good choice of framework.",
                "I used a single-page application because it is modern.",
                1,
            ),
            ("Some testing evident.", "I tested the app by clicking through it", 1),
        ],
    ),
    "sub-c": (
        {"design": 55, "implementation": 55, "testing": 55, "reflection": 62},
        "Solid work with a real user; tighten the code structure.",
        [],
    ),
    "sub-d": (
        {"design": 85, "implementation": 75, "testing": 68, "reflection": 75},
        "Very strong, ambitious project.",
        [("Excellent justification with a latency estimate.", "a latency estimate", 1)],
    ),
}


def level_label(criterion_id: str, points: int) -> str:
    name, _ = band(points)
    return MISLABELLED.get((criterion_id, points), f"{name} ({points})")


# --- Writers ------------------------------------------------------------------


def write_docx(report: dict, path: Path) -> None:
    doc = Document()
    props = doc.core_properties
    props.author = report["author"]
    props.last_modified_by = report["author"]
    props.title = report["title"]
    props.created = FIXED
    props.modified = FIXED
    doc.add_heading(report["title"], level=0)
    for heading, body in report["sections"]:
        doc.add_heading(heading, level=1)
        doc.add_paragraph(body)
    doc.save(path)
    normalise_zip(path)


# DOCX files are zip archives whose entries carry the current time. Rewrite
# them with a fixed timestamp and order so the bytes, and hashes, are stable.
ZIP_EPOCH = (2026, 1, 15, 9, 0, 0)


def normalise_zip(path: Path) -> None:
    with zipfile.ZipFile(path) as src:
        entries = [(info.filename, src.read(info)) for info in src.infolist()]
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as out:
        for name, data in sorted(entries):
            info = zipfile.ZipInfo(name, date_time=ZIP_EPOCH)
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o644 << 16
            out.writestr(info, data)


def write_pdf(report: dict, path: Path) -> None:
    styles = getSampleStyleSheet()
    story = [Paragraph(report["title"], styles["Title"])]
    for heading, body in report["sections"]:
        story += [
            Paragraph(heading, styles["Heading2"]),
            Paragraph(body, styles["BodyText"]),
            Spacer(1, 6),
        ]
    SimpleDocTemplate(
        str(path),
        pagesize=A4,
        title=report["title"],
        author=report["author"],
        invariant=1,
    ).build(story)


def write_marked_view_replica(path: Path) -> None:
    """A synthetic stand-in for a marked "current view", matching the structure
    observed in real ones: a text header page, report pages rendered as
    full-page images with digit comment markers in the margin, a text comments
    list with "Comment N | <criterion>" headings, and text rubric pages."""
    width, height = A4
    c = canvas.Canvas(str(path), pagesize=A4, invariant=1)
    c.setTitle("synthetic marked view")
    c.setFont("Helvetica-Bold", 14)
    c.drawString(60, height - 80, "Synthetic marked view (fictional)")
    c.setFont("Helvetica", 11)
    c.drawString(60, height - 105, "Grade: 62 / 100")
    c.drawString(60, height - 125, "Submission ID: 100200302")
    c.showPage()
    comments = [
        (1, "Requirements and design", "Good choice of framework."),
        (2, "Testing and evaluation", "Some testing evident."),
        (3, "Reflection and professional practice", "Descriptive rather than reflective."),
    ]
    for page_no, markers in ((1, [1]), (2, [2]), (3, [3])):
        img = Image.new("RGB", (1190, 1684), "white")
        draw = ImageDraw.Draw(img)
        for row in range(40):
            draw.text(
                (80, 80 + row * 38),
                f"Fictional report page {page_no} line {row + 1}.",
                fill="black",
            )
        buf = io.BytesIO()
        img.save(buf, format="PNG")
        buf.seek(0)
        c.drawImage(ImageReader(buf), 20, 20, width=width - 40, height=height - 40)
        for n in markers:
            c.setFillColorRGB(0.1, 0.4, 0.8)
            c.rect(width - 48, height / 2, 22, 16, fill=1, stroke=0)
            c.setFillColorRGB(1, 1, 1)
            c.setFont("Helvetica", 9)
            c.drawString(width - 42, height / 2 + 4, str(n))
        c.setFillColorRGB(0, 0, 0)
        c.showPage()
    y = height - 80
    for n, criterion, text in comments:
        c.setFont("Helvetica-Bold", 11)
        c.drawString(60, y, f"Comment {n} | {criterion}")
        c.setFont("Helvetica", 10)
        c.drawString(72, y - 18, text)
        y -= 60
    c.showPage()
    c.setFont("Helvetica", 10)
    y = height - 80
    for p in POINTS:
        name, _ = band(p)
        c.drawString(60, y, f"{name} ({p})")
        y -= 16
    c.showPage()
    c.save()


def write_rubric_csv(rubric: Rubric, path: Path) -> None:
    """The synthetic rubric as a CSV import, with one literal '\\n' artefact."""
    buf = io.StringIO()
    writer = csv.writer(buf, lineterminator="\n")
    writer.writerow(
        [
            "criterion",
            "criterion_description",
            "weight",
            "max_points",
            "level_label",
            "points",
            "descriptor",
        ]
    )
    for crit in rubric.criteria:
        description = crit.description
        if crit.id == "design":
            description = "• requirements\\n• design justification"
        for i, level in enumerate(crit.levels):
            writer.writerow(
                [
                    crit.title,
                    description if i == 0 else "",
                    f"{crit.weight:g}%" if i == 0 else "",
                    f"{crit.max_points:g}" if i == 0 else "",
                    level.label,
                    f"{level.points:g}",
                    level.descriptor,
                ]
            )
    path.write_text(buf.getvalue())


# The grid layout observed in real rubric spreadsheets: criteria down the first
# column (a title line, then bullet lines), and "Label (points)" levels across
# the first row, with a descriptor in every cell.
GRID_LEVELS = [
    ("Exceptional", 100),
    ("Excellent", 85),
    ("Very good", 75),
    ("Good", 65),
    ("Satisfactory", 55),
    ("Adequate", 45),
    ("Weak", 35),
    ("Poor", 15),
    ("None", 0),
]


def grid_rows(rubric: Rubric) -> list[list[str]]:
    rows = [[""] + [f"{label} ({points})" for label, points in GRID_LEVELS]]
    for crit in rubric.criteria:
        cell = crit.title + "\n• " + crit.description + "\n• evidence is cited"
        rows.append(
            [cell] + [f"{label}: {crit.title.lower()} (fictional)." for label, _ in GRID_LEVELS]
        )
    return rows


def write_grid_xlsx(rubric: Rubric, path: Path) -> None:
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Rubric"
    for row in grid_rows(rubric):
        ws.append(row)
    wb.properties.created = FIXED
    wb.properties.modified = FIXED
    wb.properties.creator = "synthetic"
    wb.save(path)
    normalise_zip(path)


def write_grid_docx(rubric: Rubric, path: Path) -> None:
    rows = grid_rows(rubric)
    doc = Document()
    doc.core_properties.created = FIXED
    doc.core_properties.modified = FIXED
    doc.add_heading("Synthetic rubric (fictional)", level=1)
    doc.add_paragraph("A grading scale table that is not the rubric grid:")
    scale = doc.add_table(rows=2, cols=2)
    scale.cell(0, 0).text, scale.cell(0, 1).text = "Band", "Meaning"
    scale.cell(1, 0).text, scale.cell(1, 1).text = "70–100", "First"
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    for r, row in enumerate(rows):
        for c, text in enumerate(row):
            table.cell(r, c).text = text
    doc.save(path)
    normalise_zip(path)


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def dump(path: Path, data: object) -> None:
    path.write_text(json.dumps(data, indent=2, ensure_ascii=False) + "\n")


def main() -> None:
    (PACK / "submissions").mkdir(parents=True, exist_ok=True)
    rubric = build_rubric()
    dump(PACK / "rubric.json", rubric.model_dump(mode="json"))
    write_rubric_csv(rubric, PACK / "rubric.csv")
    write_grid_xlsx(rubric, PACK / "rubric-grid.xlsx")
    write_grid_docx(rubric, PACK / "rubric-grid.docx")
    write_marked_view_replica(PACK / "marked-view-replica.pdf")

    submissions, originals, seeded = [], [], {}
    for r in REPORTS:
        path = PACK / "submissions" / f"{r['id']}.{r['format'].value}"
        (write_docx if r["format"] is SourceFormat.DOCX else write_pdf)(r, path)
        digest = sha256(path)
        submissions.append(
            Submission(
                id=r["id"],
                pseudonym=r["pseudonym"],
                source_kind=SourceKind.ORIGINAL,
                source_format=r["format"],
                source_sha256=digest,
                provenance=Provenance(
                    source=f"file:sha256:{digest}",
                    transformation=Transformation.IMPORTED,
                    actor=MODERATOR,
                    timestamp=FIXED,
                    input_hashes=[digest],
                ),
            )
        )
        seeded[r["id"]] = {"file": path.name, "metadata_author": r["author"], **r["identifiers"]}

        points, comment, notes = ORIGINAL[r["id"]]
        overall = round(sum(points.values()) / len(points))
        originals.append(
            OriginalAssessment(
                submission_id=r["id"],
                import_route=ImportRoute.TURNITIN_CURRENT_VIEW,
                criterion_marks=[
                    OriginalCriterionMark(
                        criterion_id=c,
                        level_id=f"p{p}",
                        mark=p,
                        raw_label=level_label(c, p),
                        raw_score=f"{p} / 100",
                    )
                    for c, p in points.items()
                ],
                overall_mark=overall,
                raw_overall=f"{overall} / 100",
                overall_comment=comment,
                annotations=[
                    Annotation(text=t, anchor_text=anchor, page=page) for t, anchor, page in notes
                ],
                provenance=Provenance(
                    source="synthetic Turnitin current view",
                    transformation=Transformation.IMPORTED,
                    actor=MARKER,
                    timestamp=FIXED,
                    input_hashes=[digest],
                ),
            )
        )

    dump(PACK / "original-assessments.json", [o.model_dump(mode="json") for o in originals])
    dump(PACK / "seeded-identifiers.json", seeded)

    # A small, valid example record:
    # - sub-b reviewed blind: judged first, revealed, then one criterion revised;
    # - sub-a reviewed openly, with the original marks visible throughout.
    def recorded(minutes: int) -> Provenance:
        return Provenance(
            source="moderator review",
            transformation=Transformation.RECORDED,
            actor=MODERATOR,
            timestamp=at(minutes),
        )

    blind = [
        ("design", "p48", None, "Requirements are listed but the design choice is unjustified."),
        ("implementation", "p48", "p55", "Core works; timetable bug unresolved."),
        ("testing", "p35", None, "No planned or automated testing."),
        ("reflection", "p42", None, "Descriptive rather than reflective."),
    ]
    judgements = [
        ModeratorJudgement(
            submission_id="sub-b",
            criterion_id=c,
            mode=ReviewMode.BLIND,
            first=JudgementEntry(level_id=first, comment=note, recorded_at=at(10 + i)),
            revealed_at=at(20),
            revised=(JudgementEntry(level_id=revised, recorded_at=at(25)) if revised else None),
            provenance=recorded(25),
        )
        for i, (c, first, revised, note) in enumerate(blind)
    ] + [
        ModeratorJudgement(
            submission_id="sub-a",
            criterion_id=c,
            mode=ReviewMode.OPEN,
            first=JudgementEntry(level_id=f"p{p}", recorded_at=at(35)),
            provenance=recorded(35),
        )
        for c, p in ORIGINAL["sub-a"][0].items()
    ]
    verdicts = [
        SubmissionVerdict(
            submission_id="sub-a",
            verdict=Verdict.AGREE,
            comment="Marks and comments are well supported by the work.",
            provenance=recorded(36),
        ),
        SubmissionVerdict(
            submission_id="sub-b",
            verdict=Verdict.GENEROUS,
            suggested_mark=45,
            comment="Design, implementation, and testing are marked a band or more too high.",
            provenance=recorded(26),
        ),
    ]
    record = ModerationRecord(
        id="example-moderation",
        context=ModerationContext(
            cohort_size=4,
            multiple_groups=False,
            band_distribution=[
                BandCount(label="70+", count=2),
                BandCount(label="60-69", count=1),
                BandCount(label="50-59", count=1),
            ],
            sample_note="All four submissions sampled (synthetic).",
            provenance=Provenance(
                source="synthetic moderation request",
                transformation=Transformation.ENTERED,
                actor=MODERATOR,
                timestamp=FIXED,
            ),
        ),
        rubric=rubric,
        submissions=submissions,
        original_assessments=originals,
        judgements=judgements,
        verdicts=verdicts,
        overall_comment="Marking of sub-b appears generous; sub-a is marked appropriately.",
        provenance=Provenance(
            source="synthetic fixture generator",
            transformation=Transformation.RECORDED,
            actor=MODERATOR,
            timestamp=at(40),
        ),
    )
    dump(PACK / "moderation-record.example.json", record.model_dump(mode="json"))
    print(f"wrote synthetic pack to {PACK}")


if __name__ == "__main__":
    main()
